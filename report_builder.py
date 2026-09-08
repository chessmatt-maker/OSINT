import io
import streamlit as st
import docx
from urllib.parse import urlparse
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


def add_hyperlink(paragraph, url, text, color="004B87", underline=True):
    """Inserts a clickable hyperlink into a python-docx paragraph."""
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    hyperlink.set(qn('w:history'), '1')

    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')

    if color:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), color)
        rPr.append(c)
    if underline:
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rPr.append(u)

    new_run.append(rPr)
    text_el = OxmlElement('w:t')
    text_el.text = str(text)
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

    return hyperlink


def set_cell_background(cell, hex_color):
    """Fills a table cell background with a specific hex color."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))


def create_docx_report(analysis_data, maigret_data, hibp_data, ground_truth_text, emails_found, phones_found):
    """Constructs a formal Word (.docx) document in memory."""
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Title & Header
    title = doc.add_paragraph()
    title_run = title.add_run("GGH OSINT ENTITY RESOLUTION DOSSIER")
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(17, 24, 39)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    conf_notice = doc.add_paragraph()
    conf_run = conf_notice.add_run("PRIVILEGED & CONFIDENTIAL // WORK PRODUCT")
    conf_run.font.size = Pt(9.5)
    conf_run.font.bold = True
    conf_run.font.color.rgb = RGBColor(185, 28, 28)
    conf_notice.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Ground Truth Metadata Block
    gt_header_p = doc.add_paragraph()
    gt_header_p.paragraph_format.space_after = Pt(2)
    run_gt = gt_header_p.add_run("GROUND TRUTH DATA:")
    run_gt.bold = True

    gt_content_p = doc.add_paragraph()
    gt_content_p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    gt_content_p.paragraph_format.space_after = Pt(12)

    clean_gt_text = "\n".join([line.strip() for line in ground_truth_text.splitlines() if line.strip()])
    gt_content_p.add_run(clean_gt_text)

    meta_p = doc.add_paragraph()
    
    # Display all target emails
    if emails_found:
        run_emails = meta_p.add_run("TARGET EMAILS: ")
        run_emails.bold = True
        meta_p.add_run(", ".join(emails_found) + "\n")
    
    # Display all target phones
    if phones_found:
        run_phones = meta_p.add_run("TARGET PHONES: ")
        run_phones.bold = True
        meta_p.add_run(", ".join(phones_found) + "\n")

    run_sources = meta_p.add_run("INTELLIGENCE SOURCES: ")
    run_sources.bold = True
    meta_p.add_run(
        "Have I Been Pwned (HIBP) Breach Data, Maigret Social Enumeration, "
        "Holehe Account Registration Check, Sherlock Username Enumeration, "
        "Phonenumbers Carrier & Geographic Analysis, WHOIS Domain Registration Data"
    )
    meta_p.paragraph_format.space_after = Pt(14)

    # ==========================================
    # SECTION 1: DISCOVERED SOCIAL ACCOUNTS
    # ==========================================
    h1 = doc.add_heading("1. Discovered Accounts & Online Footprint", level=1)
    h1.paragraph_format.space_before = Pt(12)

    warn_p = doc.add_paragraph()
    warn_run = warn_p.add_run(
        "WARNING: The following profiles are loosely associated based on username matching. Many results may be spam, inactive, or belong to unrelated individuals. Manual verification is required."
    )
    warn_run.font.color.rgb = RGBColor(185, 28, 28)
    warn_run.font.italic = True
    warn_run.bold = True
    warn_p.paragraph_format.space_after = Pt(8)

    # Keyword and Tag Blacklists
    adult_dating_keywords = [
        'porn', 'cam', 'xvideos', 'chaturbate', 'onlyfans', 'redtube', 'xhamster',
        'stripchat', 'tube8', 'xnxx', 'spankbang', 'nude', 'fetlife', 'escort',
        'tinder', 'bumble', 'badoo', 'okcupid', 'pof', 'plentyoffish', 'match.com',
        'hily', 'dating', 'adultfriendfinder', 'ashleymadison', 'meetme', 'zoosk'
    ]

    blacklisted_tags = {'porn', 'adult', 'nsfw', 'dating', 'cam', 'sex'}
    russian_domains = ('.ru', '.su', '.рф', '.xn--p1ai')
    russian_keywords = ['vk.com', 'vkontakte', 'ok.ru', 'odnoklassniki', 'mail.ru', 'yandex', 'rutube', 'rambler']
    
    # Useless/noise sources that produce no valuable intelligence
    useless_sources = ['fixya', 'picsart', 'joyreactor', 'pling']

    profile_records = []
    if isinstance(maigret_data, dict):
        for email_key, email_maigret_results in maigret_data.items():
            if isinstance(email_maigret_results, list):
                for entry in email_maigret_results:
                    site = entry.get("sitename") or entry.get("site") or "Unknown"
                    url = entry.get("url_user") or entry.get("url") or ""
                    tags = [t.lower() for t in entry.get("tags", [])]
                    status = entry.get("status", {}).get("status") if isinstance(entry.get("status"), dict) else entry.get(
                        "status", "Claimed")

                    if url:
                        parsed_host = urlparse(url).hostname or ""
                        parsed_host = parsed_host.lower()
                        combined_check = f"{site} {url}".lower()

                        # 0. Filter useless sources
                        is_useless = any(source in combined_check for source in useless_sources)

                        # 1. Filter Adult & Dating (Tags & URL/Site strings)
                        is_adult_or_dating = any(tag in blacklisted_tags for tag in tags) or \
                                            any(kw in combined_check for kw in adult_dating_keywords)

                        # 2. Filter Russian Sites (TLDs, platforms, hostnames)
                        is_russian = parsed_host.endswith(russian_domains) or \
                                    any(rk in combined_check for rk in russian_keywords)

                        if not is_useless and not is_adult_or_dating and not is_russian:
                            profile_records.append((site, url, status))

    if profile_records:
        table_profiles = doc.add_table(rows=1, cols=4)
        table_profiles.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_profiles.autofit = False

        headers_p = ["Platform", "Profile Link", "Status", "Verified (✓ / X)"]
        col_widths_p = [Inches(1.4), Inches(3.2), Inches(0.9), Inches(1.3)]

        hdr_p_cells = table_profiles.rows[0].cells
        for idx, text in enumerate(headers_p):
            hdr_p_cells[idx].text = text
            hdr_p_cells[idx].paragraphs[0].runs[0].font.bold = True
            hdr_p_cells[idx].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            set_cell_background(hdr_p_cells[idx], "1F2937")
            hdr_p_cells[idx].width = col_widths_p[idx]

        for site, url, status in profile_records:
            row_cells = table_profiles.add_row().cells
            for idx in range(4):
                row_cells[idx].width = col_widths_p[idx]

            row_cells[0].text = site
            row_cells[0].paragraphs[0].runs[0].font.bold = True

            p = row_cells[1].paragraphs[0]
            add_hyperlink(p, url, url)

            row_cells[2].text = str(status)
            row_cells[3].text = ""
    else:
        doc.add_paragraph("No relevant non-filtered profile links were detected or extracted.")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ==========================================
    # SECTION 2: CONFIDENCE TIERS TABLE
    # ==========================================
    h2 = doc.add_heading("2. Chain of Connection & Confidence Tiers", level=1)
    h2.paragraph_format.space_before = Pt(12)

    tier_1 = analysis_data.get("tier_1_confirmed", [])
    tier_2 = analysis_data.get("tier_2_probable", [])
    tier_3 = analysis_data.get("tier_3_unverified", [])

    tier_rows = []
    for item in tier_1:
        tier_rows.append(("Tier 1 (Confirmed)", item.get("finding", ""), item.get("reasoning", "")))
    for item in tier_2:
        tier_rows.append(("Tier 2 (Probable)", item.get("finding", ""), item.get("reasoning", "")))
    for item in tier_3:
        tier_rows.append(("Tier 3 (Unverified)", item.get("finding", ""), item.get("reasoning", "")))

    if tier_rows:
        table_tiers = doc.add_table(rows=1, cols=3)
        table_tiers.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_tiers.autofit = False

        headers = ["Confidence Tier", "Identified Finding", "Corroboration Reasoning"]
        col_widths = [Inches(1.8), Inches(2.2), Inches(2.5)]

        hdr_cells = table_tiers.rows[0].cells
        for idx, header_text in enumerate(headers):
            hdr_cells[idx].text = header_text
            hdr_cells[idx].paragraphs[0].runs[0].font.bold = True
            hdr_cells[idx].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            set_cell_background(hdr_cells[idx], "1F2937")
            hdr_cells[idx].width = col_widths[idx]

        for tier, finding, reasoning in tier_rows:
            row_cells = table_tiers.add_row().cells
            for idx in range(3):
                row_cells[idx].width = col_widths[idx]

            row_cells[0].text = tier
            row_cells[0].paragraphs[0].runs[0].font.bold = True

            if "Tier 1" in tier:
                set_cell_background(row_cells[0], "DCFCE7")
            elif "Tier 2" in tier:
                set_cell_background(row_cells[0], "FEF9C3")
            else:
                set_cell_background(row_cells[0], "F3F4F6")

            row_cells[1].text = finding
            row_cells[2].text = reasoning
    else:
        doc.add_paragraph("No specific tier associations were resolved.")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ==========================================
    # SECTION 3: TARGETED GOOGLE DORKS
    # ==========================================
    h3 = doc.add_heading("3. Targeted Google Dorks & Research Workspace", level=1)
    h3.paragraph_format.space_before = Pt(12)

    inst = doc.add_paragraph(
        "Copy the queries below into Google Search. Record any findings or relevant links directly in the workspace boxes.")
    inst.runs[0].font.italic = True

    dorks = analysis_data.get("google_dorks", [])

    for idx, dork in enumerate(dorks, 1):
        category = dork.get("category", "General Query")
        query = dork.get("query", "")
        purpose = dork.get("purpose", "")

        dp = doc.add_paragraph()
        dp.paragraph_format.space_before = Pt(8)
        dp.paragraph_format.space_after = Pt(2)
        r_num = dp.add_run(f"Query #{idx} — [{category}]\n")
        r_num.bold = True
        r_pur = dp.add_run(f"Objective: {purpose}\n")
        r_pur.font.italic = True

        r_box = dp.add_run(f"SEARCH STRING: {query}")
        r_box.bold = True
        r_box.font.color.rgb = RGBColor(14, 116, 144)

        note_table = doc.add_table(rows=1, cols=1)
        note_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = note_table.rows[0].cells[0]
        cell.width = Inches(6.5)
        set_cell_background(cell, "F9FAFB")

        np = cell.paragraphs[0]
        np.paragraph_format.space_before = Pt(4)
        np.paragraph_format.space_after = Pt(36)
        n_label = np.add_run("Findings / Notes / Case Citations:\n")
        n_label.font.size = Pt(9)
        n_label.font.color.rgb = RGBColor(156, 163, 175)

    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)
    docx_buffer.seek(0)
    return docx_buffer


def render_report(analysis_data, maigret_data, hibp_data, ground_truth_text, emails_found, phones_found):
    """Renders on-screen confirmation and provides the download button."""
    st.markdown("## 📋 Export Final Dossier")
    st.success("Investigation data cross-referenced and structured successfully.")

    docx_file = create_docx_report(
        analysis_data,
        maigret_data,
        hibp_data,
        ground_truth_text,
        emails_found,
        phones_found
    )

    # Generate filename from first email or phone
    if emails_found:
        username = emails_found[0].split('@')[0]
    elif phones_found:
        username = phones_found[0].replace('-', '').replace(' ', '')
    else:
        username = "Subject"
    
    file_name = f"OSINT_Dossier_{username}.docx"

    st.download_button(
        label="📥 Download Word Dossier (.docx)",
        data=docx_file,
        file_name=file_name,
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        use_container_width=True
    )
