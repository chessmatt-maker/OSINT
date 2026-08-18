import io
import copy
import requests
from datetime import datetime

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Mm, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

def _fetch_image_stream(url: str):
    """Download an image URL and return a BytesIO stream, or None on failure."""
    if not url:
        return None
    try:
        response = requests.get(url, timeout=3)
        if response.status_code == 200:
            return io.BytesIO(response.content)
    except Exception:
        pass
    return None


def _set_cell_text(cell, text: str):
    """Replace all runs in the first paragraph of a cell with a single plain-text run."""
    para = cell.paragraphs[0]
    for run in para.runs:
        run.text = ""
    for r in para._p.findall(qn('w:r')):
        para._p.remove(r)
    para.add_run(text or "")


def _insert_image_in_cell(cell, image_stream, width_mm: int = 25):
    """Insert an inline image into a cell's first paragraph."""
    from docx.shared import Mm as _Mm
    para = cell.paragraphs[0]
    for r in para._p.findall(qn('w:r')):
        para._p.remove(r)
    run = para.add_run()
    run.add_picture(image_stream, width=_Mm(width_mm))


def _populate_findings_table(template_tbl_elem, findings: list, is_high_confidence: bool):
    """
    For each finding, deep-copy the template table element and insert it
    *before* the template. The template is removed afterwards.
    Returns a list of new table XML elements (for the caller to wrap if needed).
    """
    from docx.table import Table

    new_tables = []
    for finding in findings:
        tbl_copy = copy.deepcopy(template_tbl_elem)
        template_tbl_elem.addprevious(tbl_copy)

        # Wrap as a python-docx Table so we can use .rows
        tbl_obj = Table(tbl_copy, template_tbl_elem.getparent())
        rows = tbl_obj.rows

        # Row 0: Thumbnail
        img_stream = _fetch_image_stream(finding.get("avatar_url", ""))
        if img_stream:
            _insert_image_in_cell(rows[0].cells[1], img_stream)
        else:
            _set_cell_text(rows[0].cells[1], "N/A")

        # Row 1: Platform (+ tier for high-confidence)
        platform = finding.get("platform", "N/A")
        if is_high_confidence:
            _set_cell_text(rows[1].cells[1],
                           f"{platform} ({finding.get('confidence_tier', '')})")
        else:
            _set_cell_text(rows[1].cells[1], platform)

        # Row 2: URL / Handle
        _set_cell_text(rows[2].cells[1], finding.get("url", "N/A"))

        # Row 3: Rationale / Flag Reason
        _set_cell_text(rows[3].cells[1], finding.get("connection_explanation", ""))

        # Row 4: Case Relevance (high-confidence only; template has 5 rows)
        if is_high_confidence and len(rows) > 4:
            _set_cell_text(rows[4].cells[1], finding.get("case_relevance", ""))

        new_tables.append(tbl_copy)

    # Remove the blueprint
    template_tbl_elem.getparent().remove(template_tbl_elem)
    return new_tables


def _replace_flat_placeholders(doc: Document, context: dict):
    """Replace simple {{ key }} placeholders in all paragraphs and table cells."""
    targets = []
    for para in doc.paragraphs:
        targets.append(para)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    targets.append(para)

    for para in targets:
        full_text = "".join(r.text or "" for r in para.runs)
        changed = False
        for key, value in context.items():
            placeholder = "{{ " + key + " }}"
            if placeholder in full_text:
                full_text = full_text.replace(placeholder, str(value))
                changed = True
            # Also handle variants without spaces
            placeholder2 = "{{" + key + "}}"
            if placeholder2 in full_text:
                full_text = full_text.replace(placeholder2, str(value))
                changed = True
        if changed:
            for run in para.runs:
                run.text = ""
            for r in para._p.findall(qn('w:r')):
                para._p.remove(r)
            para.add_run(full_text)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def generate_docx_report(ground_truth: dict, ai_analysis: dict, actionable_contact_vectors: dict = None,
                         template_path="templates/report_template.docx") -> io.BytesIO:
    """
    Generates a Word report programmatically using python-docx.

    Flat {{ variable }} placeholders in the template are resolved via simple
    string substitution.  Repeating sections (findings tables, Google Dorks)
    are populated by cloning the template table / paragraphs in Python,
    completely bypassing docxtpl's Jinja loop parser and the
    TemplateSyntaxError it produces on complex Word XML.
    """
    doc = Document(template_path)

    # ------------------------------------------------------------------
    # 1. Categorise findings
    # ------------------------------------------------------------------
    findings = ai_analysis.get("findings", [])
    high_confidence = [f for f in findings if f.get("confidence_tier") in ("Tier 1", "Tier 2")]
    dubious = [f for f in findings if f.get("confidence_tier") == "Tier 3"]

    # ------------------------------------------------------------------
    # 2. Build flat-variable context and substitute throughout the doc
    # ------------------------------------------------------------------
    acv = actionable_contact_vectors or {}
    context = {
        "report_date": datetime.now().strftime("%B %d, %Y"),
        "subject_name": ground_truth.get("primary_name", "Unknown Subject"),
        "dob": ground_truth.get("dob", "N/A"),
        "known_locations": ", ".join(ground_truth.get("locations", [])),
        "known_emails": ", ".join(ground_truth.get("emails", [])),
        "executive_summary": ai_analysis.get("executive_summary", "No summary generated."),
        "primary_phone": acv.get("primary_phone", ""),
        "primary_email": acv.get("primary_email", ""),
    }
    _replace_flat_placeholders(doc, context)

    # ------------------------------------------------------------------
    # 3. Populate findings tables
    #    Template Table 0  = header (already resolved above)
    #    Template Table 1  = high-confidence finding blueprint (5 rows x 2 cols)
    #    Template Table 2  = dubious finding blueprint (4 rows x 2 cols)
    # ------------------------------------------------------------------
    tables = doc.tables
    hc_template_elem = tables[1]._tbl
    dubious_template_elem = tables[2]._tbl

    # --- High-confidence findings (Table 1) ---
    if high_confidence:
        _populate_findings_table(hc_template_elem, high_confidence, is_high_confidence=True)
    else:
        no_findings_para = OxmlElement('w:p')
        rpr = OxmlElement('w:r')
        t_node = OxmlElement('w:t')
        t_node.text = "No high-confidence findings were identified during this scan."
        rpr.append(t_node)
        no_findings_para.append(rpr)
        hc_template_elem.addprevious(no_findings_para)
        hc_template_elem.getparent().remove(hc_template_elem)

    # --- Dubious findings (Table 2) ---
    if dubious:
        _populate_findings_table(dubious_template_elem, dubious, is_high_confidence=False)
    else:
        no_dubious_para = OxmlElement('w:p')
        rpr = OxmlElement('w:r')
        t_node = OxmlElement('w:t')
        t_node.text = "No dubious findings were identified."
        rpr.append(t_node)
        no_dubious_para.append(rpr)
        dubious_template_elem.addprevious(no_dubious_para)
        dubious_template_elem.getparent().remove(dubious_template_elem)

    # ------------------------------------------------------------------
    # 4. Populate Google Dorks
    #    The template has two paragraphs that act as a blueprint:
    #      "Targeted Query: {{ dork }}"
    #      "Investigator Notes / Screenshots: [ Paste your manual findings here... ]"
    #    We clone these for every dork, then remove the blueprint paragraphs.
    # ------------------------------------------------------------------
    google_dorks = ai_analysis.get("google_dorks", [])

    # Find the blueprint paragraphs by their text content
    body = doc.element.body
    dork_query_para = None
    dork_notes_para = None
    for para in doc.paragraphs:
        text = para.text
        if "Targeted Query:" in text and "{{ dork" in text:
            dork_query_para = para
        elif "Investigator Notes / Screenshots:" in text:
            dork_notes_para = para

    if dork_query_para is not None:
        insert_anchor = dork_query_para._element  # insert new paras before this

        if google_dorks:
            for dork in google_dorks:
                # Clone query paragraph
                query_clone = copy.deepcopy(dork_query_para._element)
                # Replace text in the clone
                for t_el in query_clone.findall('.//' + qn('w:t')):
                    if t_el.text:
                        t_el.text = t_el.text.replace("{{ dork }}", str(dork)) \
                                             .replace("{{ dork", str(dork)) \
                                             .replace("}}", "")
                insert_anchor.addprevious(query_clone)

                # Clone notes paragraph
                if dork_notes_para is not None:
                    notes_clone = copy.deepcopy(dork_notes_para._element)
                    insert_anchor.addprevious(notes_clone)

                # Blank separator paragraph
                sep = OxmlElement('w:p')
                insert_anchor.addprevious(sep)

        # Remove blueprint paragraphs
        dork_query_para._element.getparent().remove(dork_query_para._element)
        if dork_notes_para is not None:
            dork_notes_para._element.getparent().remove(dork_notes_para._element)

        if not google_dorks:
            no_dorks_para = OxmlElement('w:p')
            rpr = OxmlElement('w:r')
            t = OxmlElement('w:t')
            t.text = "No Google Dork queries were generated."
            rpr.append(t)
            no_dorks_para.append(rpr)
            body.append(no_dorks_para)

    # ------------------------------------------------------------------
    # 5. Save to BytesIO and return
    # ------------------------------------------------------------------
    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    return doc_io


def _list_value(value) -> str:
    if isinstance(value, list):
        return ", ".join([str(v) for v in value if str(v).strip()]) or "N/A"
    if value is None:
        return "N/A"
    text = str(value).strip()
    return text or "N/A"


def generate_pdf_report(
    ground_truth: dict,
    ai_analysis: dict,
    external_summary: dict,
    actionable_contact_vectors: dict = None,
) -> io.BytesIO:
    """Generate an attorney-friendly PDF report for OSINT findings."""
    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=LETTER,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        textColor=colors.HexColor("#1B365D"),
        spaceAfter=16,
    )
    section_style = ParagraphStyle(
        "Section",
        parent=styles["Heading2"],
        textColor=colors.HexColor("#1B365D"),
        spaceBefore=10,
        spaceAfter=8,
    )
    body_style = styles["BodyText"]

    acv = actionable_contact_vectors or {}
    confirmed = (external_summary or {}).get("confirmed", [])
    probable = (external_summary or {}).get("probable", [])
    unrelated = (external_summary or {}).get("unrelated", [])
    counts = (external_summary or {}).get("counts", {})

    story = [
        Paragraph("OSINT Identity & Exposure Analysis Report", title_style),
        Paragraph(f"Report Date: {datetime.now().strftime('%B %d, %Y')}", body_style),
        Spacer(1, 0.12 * inch),
    ]

    overview_rows = [
        ["Subject Name", _list_value(ground_truth.get("primary_name", ground_truth.get("name", "N/A")))],
        ["Known Aliases", _list_value(ground_truth.get("aliases", []))],
        ["Known Locations", _list_value(ground_truth.get("locations", []))],
        ["Known Emails", _list_value(ground_truth.get("emails", []))],
        ["Known Usernames", _list_value(ground_truth.get("usernames", []))],
        ["Primary Verified Phone", _list_value(acv.get("primary_phone", ""))],
        ["Primary Verified Email", _list_value(acv.get("primary_email", ""))],
    ]
    overview_table = Table(overview_rows, colWidths=[1.9 * inch, 5.3 * inch])
    overview_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF3FA")),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.black),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#C6D3E5")),
        ("ROWBACKGROUNDS", (1, 0), (1, -1), [colors.white, colors.HexColor("#FAFCFF")]),
    ]))
    story.extend([Paragraph("Subject Baseline", section_style), overview_table, Spacer(1, 0.16 * inch)])

    story.append(Paragraph("Executive Summary", section_style))
    story.append(Paragraph(ai_analysis.get("executive_summary", "No executive summary generated."), body_style))
    story.append(Spacer(1, 0.16 * inch))

    summary_rows = [
        ["Confirmed Matches", str(counts.get("confirmed", 0))],
        ["Probable Matches", str(counts.get("probable", 0))],
        ["Unrelated/Noise", str(counts.get("unrelated", 0))],
    ]
    summary_table = Table(summary_rows, colWidths=[2.4 * inch, 1.2 * inch])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1B365D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#C6D3E5")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F7FD")]),
    ]))
    story.extend([Paragraph("Match Overview", section_style), summary_table, Spacer(1, 0.16 * inch)])

    def _add_record_section(title: str, records: list):
        story.append(Paragraph(title, section_style))
        if not records:
            story.append(Paragraph("No entries.", body_style))
            story.append(Spacer(1, 0.1 * inch))
            return
        for idx, rec in enumerate(records, start=1):
            detail = rec.get("detail", "")
            source = rec.get("source", "Unknown")
            story.append(Paragraph(f"<b>{idx}. [{source}]</b> {detail}", body_style))
            extras = []
            for key in ("platform", "url", "email", "username", "breach_name", "database_name"):
                value = rec.get(key)
                if value:
                    extras.append(f"{key.replace('_', ' ').title()}: {value}")
            if extras:
                story.append(Paragraph(" • " + " | ".join(extras), body_style))
            story.append(Spacer(1, 0.06 * inch))

    _add_record_section("Confirmed Public-Facing Accounts / Exposure Links", confirmed)
    _add_record_section("Probable Matches Requiring Manual Verification", probable)
    _add_record_section("Unrelated or Low-Confidence Noise", unrelated)

    story.append(Paragraph("Attorney Notes", section_style))
    story.append(Paragraph(
        "The findings above are organized by confidence. Confirmed entries align with known identifiers from the "
        "ground-truth baseline, while probable entries require corroboration before legal reliance.",
        body_style,
    ))

    doc.build(story)
    output.seek(0)
    return output